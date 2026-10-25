from pathlib import Path
import hashlib
import uuid
import pandas as pd
import fitz
import pytesseract
from PIL import Image
from docx import Document as WordDocument
from config.settings import UPLOAD_PATH, TESSERACT_CMD, OCR_LANGUAGE

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm", ".csv"}

if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD

def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def save_upload(uploaded_file):
    data = uploaded_file.getvalue()
    digest = sha256_bytes(data)
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported format: {suffix or '(no extension)'}")
    doc_id = str(uuid.uuid4())
    safe_name = Path(uploaded_file.name).name
    stored_path = UPLOAD_PATH / f"{doc_id}{suffix}"
    stored_path.write_bytes(data)
    return {
        "id": doc_id, "filename": safe_name, "stored_path": str(stored_path),
        "file_hash": digest, "file_type": suffix.lstrip("."),
        "file_size": len(data), "status": "uploaded"
    }

def _extract_pdf(path, enable_ocr=True, ocr_min_chars=30):
    pages = []
    with fitz.open(path) as pdf:
        for index, page in enumerate(pdf):
            text = page.get_text("text").strip()
            if enable_ocr and len(text) < ocr_min_chars:
                try:
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
                    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_text = pytesseract.image_to_string(image, lang=OCR_LANGUAGE).strip()
                    if len(ocr_text) > len(text):
                        text = ocr_text
                except pytesseract.pytesseract.TesseractNotFoundError:
                    raise RuntimeError(
                        "Tesseract was not found. Install Tesseract OCR or set TESSERACT_CMD in .env."
                    )
            if text:
                pages.append({"text": text, "source": f"page {index + 1}", "page": index + 1})
    return pages

def _extract_docx(path):
    doc = WordDocument(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table_index, table in enumerate(doc.tables, start=1):
        for row in table.rows:
            vals = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            if any(vals):
                parts.append(f"Table {table_index}: " + " | ".join(vals))
    text = "\n".join(parts)
    return [{"text": text, "source": "document", "page": None}] if text else []

def _extract_spreadsheet(path, is_csv=False):
    sheets = {}
    if is_csv:
        sheets["CSV"] = pd.read_csv(path, dtype=str, keep_default_na=False, encoding_errors="replace")
    else:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False, engine="openpyxl")
    pages = []
    for sheet_name, frame in sheets.items():
        frame = frame.fillna("")
        columns = [str(c) for c in frame.columns]
        for start in range(0, len(frame), 30):
            subset = frame.iloc[start:start + 30]
            lines = [f"Columns: {' | '.join(columns)}"]
            for row_index, (_, row) in enumerate(subset.iterrows(), start=start + 2):
                values = [f"{columns[i]}: {str(row.iloc[i])}" for i in range(len(columns))
                          if str(row.iloc[i]).strip() and str(row.iloc[i]).lower() != "nan"]
                if values:
                    lines.append(f"Row {row_index}: " + " | ".join(values))
            if len(lines) > 1:
                pages.append({
                    "text": f"Sheet: {sheet_name}\n" + "\n".join(lines),
                    "source": f"sheet {sheet_name}, rows {start + 2}-{start + len(subset) + 1}",
                    "page": None,
                })
    return pages

def extract_document(path, enable_ocr=True):
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, enable_ocr=enable_ocr)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        text = Path(path).read_text(encoding="utf-8", errors="replace").strip()
        return [{"text": text, "source": "document", "page": None}] if text else []
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_spreadsheet(path)
    if suffix == ".csv":
        return _extract_spreadsheet(path, is_csv=True)
    raise ValueError(f"Unsupported file format: {suffix}")
